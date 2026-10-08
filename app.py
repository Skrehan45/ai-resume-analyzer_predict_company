import streamlit as st
import io
import os
import re
import html
import shutil

import fitz  # PyMuPDF  ->  pip install pymupdf
import pytesseract
from PIL import Image
from docx import Document


# =========================================================
# TESSERACT SETUP (Windows fallback path)
# =========================================================

_WIN_TESSERACT = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

if not shutil.which("tesseract") and os.path.exists(_WIN_TESSERACT):
    pytesseract.pytesseract.tesseract_cmd = _WIN_TESSERACT


# =========================================================
# PAGE CONFIG
# =========================================================

st.set_page_config(
    page_title="AI Resume Analyzer",
    page_icon="🤖",
    layout="wide",
    initial_sidebar_state="collapsed"
)


# =========================================================
# CSS
# =========================================================

st.markdown("""
<style>

.stApp {
    background: linear-gradient(135deg, #080d1c, #0d1630, #080d1c);
    color: white;
}

.block-container {
    max-width: 1150px;
    padding-top: 25px;
    padding-bottom: 50px;
}

.hero {
    padding: 28px;
    border-radius: 22px;
    background: linear-gradient(135deg, #111d3a, #17264b);
    border: 1px solid #304a80;
    margin-bottom: 22px;
}

.hero h1 {
    margin: 0;
    color: white;
    font-size: 32px;
}

.hero p {
    color: #a8b8dc;
}

.step {
    text-align: center;
    padding: 14px 6px;
    border-radius: 15px;
    background: #101a31;
    border: 1px solid #293c68;
    color: #91a2c4;
    min-height: 70px;
}

.step.active {
    background: linear-gradient(135deg, #315ee8, #7047e8);
    color: white;
    border-color: #7d98ff;
}

.step-num {
    font-weight: 800;
    font-size: 18px;
}

.step-title {
    font-size: 12px;
    margin-top: 4px;
}

.card {
    background: linear-gradient(145deg, #111a31, #0c1428);
    border: 1px solid #293d69;
    border-radius: 18px;
    padding: 20px;
    margin-bottom: 16px;
    word-wrap: break-word;
}

.metric {
    background: #101a31;
    border: 1px solid #2c4272;
    border-radius: 16px;
    padding: 20px;
    text-align: center;
    min-height: 105px;
}

.metric-label {
    color: #91a4cb;
    font-size: 13px;
}

.metric-value {
    color: white;
    font-size: 25px;
    font-weight: 800;
    margin-top: 8px;
    word-wrap: break-word;
}

.skill {
    display: inline-block;
    padding: 7px 12px;
    margin: 4px;
    border-radius: 18px;
    background: #1b315c;
    border: 1px solid #4b6daa;
    color: #d5e1ff;
    font-size: 13px;
}

.job-card {
    background: #101a31;
    border: 1px solid #2c4272;
    border-radius: 16px;
    padding: 17px;
    margin-bottom: 12px;
}

.company-card {
    background: #101a31;
    border: 1px solid #2c4272;
    border-radius: 16px;
    padding: 17px;
    margin-bottom: 12px;
}

.match {
    color: #55e6a5;
    font-weight: 800;
}

.gap {
    color: #ffcf75;
}

.info-box {
    background: #101b34;
    border-left: 4px solid #5d83ff;
    border-radius: 10px;
    padding: 13px;
    margin: 10px 0;
}

.company-name {
    color: white;
    font-size: 18px;
    font-weight: 800;
}

.company-type {
    color: #94a8cf;
    font-size: 13px;
    margin-top: 5px;
}

</style>
""", unsafe_allow_html=True)


# =========================================================
# SKILLS  (each skill appears in exactly one category)
# =========================================================

SKILLS = {
    "Programming": [
        "Python", "C", "C++", "Java", "JavaScript", "R"
    ],

    "Data Science": [
        "Pandas", "NumPy", "Matplotlib", "Seaborn",
        "Scikit-learn", "TensorFlow", "PyTorch"
    ],

    "Database": [
        "SQL", "MySQL", "PostgreSQL", "MongoDB"
    ],

    "Web Development": [
        "HTML", "CSS", "React", "Node.js", "Django", "Flask"
    ],

    "Data Analytics": [
        "Excel", "Power BI", "Tableau"
    ],

    "Cloud": [
        "AWS", "Azure", "Google Cloud"
    ],

    "Tools": [
        "Git", "GitHub", "Docker", "Jupyter", "VS Code"
    ],

    "AI / ML": [
        "Machine Learning", "Deep Learning",
        "Natural Language Processing",
        "Artificial Intelligence"
    ],

    "Soft Skills": [
        "Communication", "Leadership", "Teamwork",
        "Problem Solving", "Critical Thinking",
        "Time Management"
    ]
}


# Alternative spellings that count as the same skill
SKILL_ALIASES = {
    "Google Cloud": ["GCP", "Google Cloud Platform"],
    "Natural Language Processing": ["NLP"],
    "Artificial Intelligence": ["AI"],
    "Machine Learning": ["ML"],
    "Deep Learning": ["DL"],
    "Scikit-learn": ["sklearn", "scikit learn"],
    "Power BI": ["PowerBI"],
    "Node.js": ["NodeJS", "Node js"],
    "React": ["ReactJS", "React.js"],
    "PostgreSQL": ["Postgres"],
    "MongoDB": ["Mongo DB"],
    "TensorFlow": ["Tensor Flow"],
    "JavaScript": ["JS"],
    "AWS": ["Amazon Web Services"],
}

# Words that are also normal English words -> require proper capitalisation
CASE_SENSITIVE_SKILLS = {"Excel", "React"}

# Context words used to confirm single-letter skills such as "C" and "R"
SHORT_SKILL_CONTEXT = re.compile(
    r"language|skill|programming|technolog|tools|proficien|familiar|stack"
    r"|\bpython\b|\bjava\b|\bjavascript\b|\bsql\b|c\+\+|\bmatlab\b",
    re.IGNORECASE
)


# =========================================================
# JOB ROLES
# =========================================================

JOB_ROLES = {
    "Data Analyst": [
        "Python", "SQL", "Excel", "Power BI",
        "Pandas", "NumPy", "Matplotlib"
    ],

    "Data Scientist": [
        "Python", "SQL", "Pandas", "NumPy",
        "Scikit-learn", "Machine Learning", "Matplotlib"
    ],

    "Machine Learning Engineer": [
        "Python", "Machine Learning", "Scikit-learn",
        "TensorFlow", "PyTorch", "SQL", "Git"
    ],

    "AI Engineer": [
        "Python", "Artificial Intelligence",
        "Machine Learning", "Deep Learning",
        "TensorFlow", "PyTorch"
    ],

    "Python Developer": [
        "Python", "Django", "Flask",
        "SQL", "Git", "GitHub"
    ],

    "Web Developer": [
        "HTML", "CSS", "JavaScript",
        "React", "Node.js", "Git"
    ],

    "Backend Developer": [
        "Python", "Django", "Flask",
        "SQL", "MongoDB", "Git"
    ],

    "Business Intelligence Analyst": [
        "Excel", "Power BI", "SQL",
        "Tableau", "Python"
    ],

    "Data Engineer": [
        "Python", "SQL", "MongoDB",
        "PostgreSQL", "AWS", "Docker", "Git"
    ],

    "Cloud Engineer": [
        "AWS", "Azure", "Google Cloud",
        "Docker", "Git", "Python"
    ]
}


# =========================================================
# COMPANIES
# =========================================================

COMPANIES = {
    "Data Analyst": [
        ("Accenture", "IT Services & Consulting"),
        ("Deloitte", "Consulting & Analytics"),
        ("TCS", "IT Services"),
        ("Infosys", "IT Services & Consulting"),
        ("Capgemini", "Technology & Consulting"),
        ("Wipro", "IT Services")
    ],

    "Data Scientist": [
        ("Microsoft", "Technology & AI"),
        ("Amazon", "Technology & Cloud"),
        ("IBM", "AI & Data Solutions"),
        ("Google", "Technology & AI"),
        ("Accenture", "Analytics & Consulting"),
        ("Deloitte", "Analytics & Consulting")
    ],

    "Machine Learning Engineer": [
        ("Google", "AI & Machine Learning"),
        ("Microsoft", "AI & Cloud"),
        ("Amazon", "AI & AWS"),
        ("NVIDIA", "AI & Computing"),
        ("IBM", "AI & Enterprise"),
        ("Accenture", "AI & Consulting")
    ],

    "AI Engineer": [
        ("Google", "Artificial Intelligence"),
        ("Microsoft", "Artificial Intelligence"),
        ("NVIDIA", "AI & Computing"),
        ("Amazon", "AI & Cloud"),
        ("IBM", "AI Solutions"),
        ("Adobe", "AI & Technology")
    ],

    "Python Developer": [
        ("TCS", "Software Development"),
        ("Infosys", "Software Development"),
        ("Accenture", "Technology Services"),
        ("Wipro", "Software & IT"),
        ("Deloitte", "Technology Consulting"),
        ("Zoho", "Software Products")
    ],

    "Web Developer": [
        ("Accenture", "Web & Technology"),
        ("TCS", "Software Development"),
        ("Infosys", "Technology Services"),
        ("Wipro", "IT Services"),
        ("Cognizant", "Technology Services"),
        ("Zoho", "Software Products")
    ],

    "Backend Developer": [
        ("Amazon", "Software & Cloud"),
        ("Microsoft", "Software Development"),
        ("Google", "Software Engineering"),
        ("TCS", "IT Services"),
        ("Infosys", "Software Development"),
        ("Cognizant", "Technology Services")
    ],

    "Business Intelligence Analyst": [
        ("Deloitte", "Business Intelligence"),
        ("Accenture", "Analytics"),
        ("TCS", "Data & Analytics"),
        ("Infosys", "Analytics"),
        ("Capgemini", "Data & BI"),
        ("EY", "Analytics & Consulting")
    ],

    "Data Engineer": [
        ("Amazon", "Data & Cloud"),
        ("Microsoft", "Data & Azure"),
        ("Google", "Data & Cloud"),
        ("IBM", "Data Engineering"),
        ("Accenture", "Data Engineering"),
        ("Deloitte", "Data & Analytics")
    ],

    "Cloud Engineer": [
        ("Amazon", "Cloud / AWS"),
        ("Microsoft", "Cloud / Azure"),
        ("Google", "Cloud / GCP"),
        ("IBM", "Cloud Computing"),
        ("Accenture", "Cloud Consulting"),
        ("TCS", "Cloud Services")
    ]
}


# =========================================================
# SESSION STATE
# =========================================================

if "page" not in st.session_state:
    st.session_state.page = 1

if "resume_text" not in st.session_state:
    st.session_state.resume_text = ""

if "source_name" not in st.session_state:
    st.session_state.source_name = ""


# =========================================================
# TEXT EXTRACTION
# =========================================================

@st.cache_data(show_spinner=False)
def _read_pdf(file_bytes):
    """Read a PDF. Falls back to OCR for pages that have no embedded text."""

    pdf = fitz.open(stream=file_bytes, filetype="pdf")
    text = ""

    try:
        for page in pdf:
            page_text = page.get_text()

            if not page_text.strip():
                pix = page.get_pixmap(dpi=200)
                img = Image.open(io.BytesIO(pix.tobytes("png")))
                page_text = pytesseract.image_to_string(img)

            text += page_text + "\n"
    finally:
        pdf.close()

    return text


@st.cache_data(show_spinner=False)
def _read_docx(file_bytes):

    document = Document(io.BytesIO(file_bytes))
    text = ""

    for paragraph in document.paragraphs:
        text += paragraph.text + "\n"

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text for cell in row.cells]
            text += " ".join(cells) + "\n"

    return text


@st.cache_data(show_spinner=False)
def _read_image(image_bytes):

    image = Image.open(io.BytesIO(image_bytes)).convert("RGB")
    return pytesseract.image_to_string(image)


def extract_pdf_text(uploaded_file):
    try:
        return _read_pdf(uploaded_file.getvalue())
    except Exception as e:
        st.error("PDF Error: " + str(e))
        return ""


def extract_docx_text(uploaded_file):
    try:
        return _read_docx(uploaded_file.getvalue())
    except Exception as e:
        st.error("DOCX Error: " + str(e))
        return ""


def extract_image_text(image_file):
    try:
        return _read_image(image_file.getvalue())
    except Exception as e:
        st.error(
            "OCR Error: " + str(e)
            + " (make sure Tesseract OCR is installed)"
        )
        return ""


# =========================================================
# PERSONAL INFO
# =========================================================

def extract_name(text):

    excluded = {
        "resume", "curriculum", "vitae", "cv", "email", "phone",
        "contact", "linkedin", "github", "mobile", "address"
    }

    lines = []

    for line in text.splitlines():
        line = line.strip()
        if line:
            lines.append(line)

    for line in lines[:15]:

        words = set(re.findall(r"[a-z]+", line.lower()))

        if words & excluded:
            continue

        if "@" in line:
            continue

        if any(character.isdigit() for character in line):
            continue

        if 1 <= len(line.split()) <= 5:
            if re.fullmatch(r"[A-Za-z .'-]+", line):
                return line

    return "Not detected"


def extract_email(text):

    pattern = (
        r"[A-Za-z0-9._%+-]+"
        r"@[A-Za-z0-9.-]+"
        r"\.[A-Za-z]{2,}"
    )

    result = re.findall(pattern, text)

    if result:
        return result[0]

    return "Not detected"


def extract_phone(text):

    patterns = [
        r"(?<!\d)(?:\+91[\s-]?)?[6-9]\d{9}(?!\d)",
        r"(?<!\d)\+\d{1,3}[\s-]?\d{10}(?!\d)"
    ]

    for pattern in patterns:

        match = re.search(pattern, text)

        if match:
            return match.group(0)

    return "Not detected"


# =========================================================
# SKILL DETECTION
# =========================================================

def _term_matches(text, term):
    """Check if one term (a skill name or alias) appears in the text."""

    single_letter = len(term) == 1

    case_sensitive = (
        single_letter
        or term in CASE_SENSITIVE_SKILLS
        or (len(term) <= 3 and term.isupper())
    )

    flags = 0 if case_sensitive else re.IGNORECASE

    if single_letter:
        head = r"(?<![\w+#&])"
        tail = r"(?![\w+#&])"
    else:
        head = r"(?<!\w)"
        tail = r"(?!\w)" if term[-1] in "+#" else r"(?![\w+#])"

    variants = [term]

    if term in CASE_SENSITIVE_SKILLS:
        variants.append(term.upper())

    for variant in variants:

        pattern = head + re.escape(variant) + tail

        if single_letter:
            # "C" / "R" only count inside a skills-like line
            for line in text.splitlines():
                if (
                    re.search(pattern, line, flags)
                    and SHORT_SKILL_CONTEXT.search(line)
                ):
                    return True
        else:
            if re.search(pattern, text, flags):
                return True

    return False


def skill_exists(text, skill):

    names = [skill] + SKILL_ALIASES.get(skill, [])

    return any(_term_matches(text, name) for name in names)


def detect_skills(text):

    detected = {}

    for category, skills in SKILLS.items():

        found = []

        for skill in skills:
            if skill_exists(text, skill):
                found.append(skill)

        if found:
            detected[category] = found

    return detected


def unique_skill_set(detected_skills):

    result = set()

    for skills in detected_skills.values():
        for skill in skills:
            result.add(skill.lower())

    return result


# =========================================================
# ANALYSIS
# =========================================================

def recommend_jobs(detected_skills):

    skill_set = unique_skill_set(detected_skills)

    results = []

    for job, required in JOB_ROLES.items():

        matched = []
        missing = []

        for skill in required:
            if skill.lower() in skill_set:
                matched.append(skill)
            else:
                missing.append(skill)

        score = round(len(matched) / len(required) * 100, 1)

        results.append({
            "Job": job,
            "Score": score,
            "Matched": matched,
            "Missing": missing
        })

    results.sort(key=lambda item: item["Score"], reverse=True)

    return results


def calculate_resume_score(text, detected_skills):

    score = 0
    lower_text = text.lower()

    skill_count = len(unique_skill_set(detected_skills))

    score += min(skill_count * 3, 30)

    education_words = [
        "education", "b.tech", "btech", "bachelor", "degree",
        "university", "college", "engineering", "master"
    ]

    if any(word in lower_text for word in education_words):
        score += 15

    if "project" in lower_text:
        score += 15

    if "experience" in lower_text or "internship" in lower_text:
        score += 15

    if "certification" in lower_text or "certificate" in lower_text:
        score += 10

    if extract_email(text) != "Not detected":
        score += 5

    if extract_phone(text) != "Not detected":
        score += 5

    return min(score, 100)


# =========================================================
# UI HELPERS
# =========================================================

def show_metric(label, value):

    st.markdown(
        '<div class="metric">'
        '<div class="metric-label">'
        + html.escape(str(label))
        + '</div>'
        '<div class="metric-value">'
        + html.escape(str(value))
        + '</div>'
        '</div>',
        unsafe_allow_html=True
    )


def card_header(title, description):

    st.markdown(
        '<div class="card">'
        '<h2>' + title + '</h2>'
        '<p>' + description + '</p>'
        '</div>',
        unsafe_allow_html=True
    )


def skill_chips(skills):

    chips = ""

    for skill in skills:
        chips += '<span class="skill">' + html.escape(skill) + '</span>'

    st.markdown(chips, unsafe_allow_html=True)


# =========================================================
# HEADER
# =========================================================

st.markdown("""
<div class="hero">
<h1>🤖 AI Resume Analyzer & Job Recommendation System</h1>
<p>
Upload or scan your resume • Analyze Skills • Find Jobs •
Discover Companies • Identify Skill Gaps
</p>
</div>
""", unsafe_allow_html=True)


# =========================================================
# STEP BAR
# =========================================================

step_names = [
    "📤 Resume Input",
    "👤 Resume Overview",
    "💼 Job Matches",
    "🎯 Skill Gap"
]

step_columns = st.columns(4)

for number, column in enumerate(step_columns, start=1):

    with column:

        if number == st.session_state.page:
            css_class = "step active"
        else:
            css_class = "step"

        st.markdown(
            '<div class="' + css_class + '">'
            '<div class="step-num">' + str(number) + '</div>'
            '<div class="step-title">' + step_names[number - 1] + '</div>'
            '</div>',
            unsafe_allow_html=True
        )


st.markdown("<br>", unsafe_allow_html=True)


# =========================================================
# PAGE 1 - RESUME INPUT
# =========================================================

if st.session_state.page == 1:

    card_header(
        "📤 Resume Input",
        "Upload your resume as PDF/DOCX or scan it using your camera."
    )

    left, right = st.columns(2)

    with left:

        st.subheader("📁 Upload Resume")

        uploaded_file = st.file_uploader(
            "Choose PDF or DOCX",
            type=["pdf", "docx"]
        )

        if uploaded_file is not None:

            with st.spinner("Reading resume..."):

                if uploaded_file.name.lower().endswith(".pdf"):
                    extracted_text = extract_pdf_text(uploaded_file)
                else:
                    extracted_text = extract_docx_text(uploaded_file)

            if extracted_text.strip():

                st.session_state.resume_text = extracted_text
                st.session_state.source_name = uploaded_file.name

                st.success("✅ Resume uploaded successfully.")

            else:

                st.warning("⚠️ No readable text found.")

    with right:

        st.subheader("📷 Scan Resume")

        camera_image = st.camera_input("Take a photo of your resume")

        if camera_image is not None:

            with st.spinner("Reading resume using OCR..."):
                extracted_text = extract_image_text(camera_image)

            if extracted_text.strip():

                st.session_state.resume_text = extracted_text
                st.session_state.source_name = "Camera Scan"

                st.success("✅ Resume scanned successfully.")

            else:

                st.warning("⚠️ Could not read text from image.")

    if st.session_state.resume_text:

        st.markdown("### 📄 Resume Loaded")

        st.info("Source: " + st.session_state.source_name)

        with st.expander("Preview Extracted Resume Text"):
            st.text(st.session_state.resume_text[:5000])

        if st.button("🗑️ Clear Resume"):
            st.session_state.resume_text = ""
            st.session_state.source_name = ""
            st.rerun()


# =========================================================
# PAGE 2 - RESUME OVERVIEW
# =========================================================

elif st.session_state.page == 2:

    text = st.session_state.resume_text

    if not text.strip():

        st.warning("⚠️ Please upload or scan a resume first.")

    else:

        detected_skills = detect_skills(text)

        score = calculate_resume_score(text, detected_skills)

        name = extract_name(text)
        email = extract_email(text)
        phone = extract_phone(text)

        card_header(
            "👤 Resume Overview",
            "Information automatically detected from your resume."
        )

        c1, c2, c3 = st.columns(3)

        with c1:
            show_metric("Resume Score", str(score) + "/100")

        with c2:
            show_metric(
                "Skills Detected",
                len(unique_skill_set(detected_skills))
            )

        with c3:
            show_metric("Source", st.session_state.source_name)

        st.markdown("### 👤 Personal Information")

        p1, p2, p3 = st.columns(3)

        with p1:
            st.markdown(
                '<div class="card"><b>Name</b><br><br>'
                + html.escape(name)
                + '</div>',
                unsafe_allow_html=True
            )

        with p2:
            st.markdown(
                '<div class="card"><b>Email</b><br><br>'
                + html.escape(email)
                + '</div>',
                unsafe_allow_html=True
            )

        with p3:
            st.markdown(
                '<div class="card"><b>Phone</b><br><br>'
                + html.escape(phone)
                + '</div>',
                unsafe_allow_html=True
            )

        st.markdown("### 🛠️ Detected Skills")

        if detected_skills:

            for category, skills in detected_skills.items():

                st.markdown("#### " + category)
                skill_chips(skills)

        else:

            st.warning("No predefined skills were detected.")


# =========================================================
# PAGE 3 - JOB MATCHES
# =========================================================

elif st.session_state.page == 3:

    text = st.session_state.resume_text

    if not text.strip():

        st.warning("⚠️ Please upload a resume first.")

    else:

        detected_skills = detect_skills(text)

        results = recommend_jobs(detected_skills)

        card_header(
            "💼 Job & Company Recommendations",
            "Jobs are ranked according to skills detected in your resume."
        )

        st.markdown("### 💼 Recommended Job Roles")

        for result in results[:7]:

            job = result["Job"]
            score = result["Score"]

            matched = ", ".join(result["Matched"]) or "None"
            missing = ", ".join(result["Missing"]) or "None"

            st.markdown(
                '<div class="job-card">'
                '<h3>💼 ' + html.escape(job) + '</h3>'
                '<p class="match">Match Score: ' + str(score) + '%</p>'
                '<p><b>Matched Skills:</b> ' + html.escape(matched) + '</p>'
                '<p class="gap"><b>Missing Skills:</b> '
                + html.escape(missing) + '</p>'
                '</div>',
                unsafe_allow_html=True
            )

        st.markdown("### 🏢 Recommended Companies")

        shown_companies = set()

        for result in results[:3]:

            job = result["Job"]

            for company, company_type in COMPANIES.get(job, []):

                if company in shown_companies:
                    continue

                shown_companies.add(company)

                st.markdown(
                    '<div class="company-card">'
                    '<div class="company-name">🏢 '
                    + html.escape(company) + '</div>'
                    '<div class="company-type">'
                    + html.escape(company_type) + '</div>'
                    '<div class="company-type">Suitable role: '
                    + html.escape(job) + '</div>'
                    '</div>',
                    unsafe_allow_html=True
                )

                if len(shown_companies) >= 10:
                    break

            if len(shown_companies) >= 10:
                break

        st.info(
            "💡 Company recommendations are career-role "
            "suggestions, not live job vacancies."
        )


# =========================================================
# PAGE 4 - SKILL GAP
# =========================================================

elif st.session_state.page == 4:

    text = st.session_state.resume_text

    if not text.strip():

        st.warning("⚠️ Please upload a resume first.")

    else:

        detected_skills = detect_skills(text)

        card_header(
            "🎯 Skill Gap Analysis",
            "Select a target career role to identify the skills "
            "you should learn next."
        )

        selected_job = st.selectbox(
            "🎯 Select Target Job Role",
            list(JOB_ROLES.keys())
        )

        required_skills = JOB_ROLES[selected_job]

        detected_lower = unique_skill_set(detected_skills)

        matched = []
        missing = []

        for skill in required_skills:
            if skill.lower() in detected_lower:
                matched.append(skill)
            else:
                missing.append(skill)

        match_percentage = round(
            len(matched) / len(required_skills) * 100,
            1
        )

        c1, c2, c3 = st.columns(3)

        with c1:
            show_metric("Target Role", selected_job)

        with c2:
            show_metric("Skill Match", str(match_percentage) + "%")

        with c3:
            show_metric("Skills To Learn", len(missing))

        st.markdown("### ✅ Skills You Already Have")

        if matched:
            skill_chips(matched)
        else:
            st.info("No matching skills found for this role.")

        st.markdown("### 📚 Skills You Should Learn")

        if missing:

            for index, skill in enumerate(missing, start=1):

                st.markdown(
                    '<div class="info-box">'
                    '<b>Step ' + str(index) + '</b> → Learn <b>'
                    + html.escape(skill) + '</b>'
                    '</div>',
                    unsafe_allow_html=True
                )

        else:

            st.success("🎉 Your resume already covers all required skills.")

        st.markdown("### 🚀 Suggested Learning Roadmap")

        if missing:

            st.write(
                "Learn the missing skills, build projects "
                "using them, add the projects to your resume, "
                "and then apply for suitable roles."
            )

        else:

            st.success(
                "🎉 You are ready to start applying "
                "for this target role."
            )


# =========================================================
# NAVIGATION
# =========================================================

st.markdown("<br>", unsafe_allow_html=True)

back_col, next_col, end_col = st.columns(3)

has_resume = bool(st.session_state.resume_text.strip())


with back_col:

    if st.session_state.page > 1:

        if st.button("← Back", use_container_width=True):
            st.session_state.page -= 1
            st.rerun()


with next_col:

    if st.session_state.page < 4:

        if st.button(
            "Next →",
            use_container_width=True,
            disabled=not has_resume
        ):
            st.session_state.page += 1
            st.rerun()

        if not has_resume:
            st.caption("Upload or scan a resume to continue.")


with end_col:

    if st.session_state.page == 4:

        if st.button("✓ Analysis Complete", use_container_width=True):
            st.balloons()
            st.success("🎉 Resume analysis completed successfully!")
